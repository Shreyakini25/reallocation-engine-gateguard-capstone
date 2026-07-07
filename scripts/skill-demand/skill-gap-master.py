#!/usr/bin/env python3
"""
skill-gap-master.py — Master skill gap scanner for DE/DA roles at verified H-1B sponsors.

Part of: case-de-da-live-skill-gap recipe

What it does:
  1. Reads my_targets.txt (one company name per line)
  2. Looks up each company in the 80-days CSV for H-1B history
  3. Hits Greenhouse then Lever public APIs for live DE/DA job titles
  4. Extracts skills from live titles and counts frequency
  5. Writes skill_gap_report.xlsx with 4 sheets:
       Sheet 1 — Skill Rankings     : ranked skills with cognitive tier
       Sheet 2 — Sponsor Scorecard  : per-company H-1B + live role data
       Sheet 3 — Live Job Titles    : every DE/DA title found
       Sheet 4 — Not Found          : companies with no Greenhouse/Lever presence

Usage:
  python3 scripts/skill-demand/skill-gap-master.py
  python3 scripts/skill-demand/skill-gap-master.py --targets my_targets.txt
  python3 scripts/skill-demand/skill-gap-master.py --dry-run
  python3 scripts/skill-demand/skill-gap-master.py --all-sponsors   # ignore targets file, scan all 38 sponsors

Input:
  my_targets.txt — one company name per line (place in repo root)
  Example:
    Databricks
    CVS Health
    Experian
    Cotiviti

Output:
  data/skill-demand/skill_gap_report.xlsx

Prime directive (SNICKERDOODLE.md):
  Verified local data first. LLM judgment not used anywhere in this script.
  All skill and cognitive tier classifications are from a static human-curated mapping.

Author: Komal Pravinkumar
"""

from __future__ import annotations

import argparse
import re
import sys
import time
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Import core extraction engine from jd-skill-extractor
sys.path.insert(0, str(Path(__file__).resolve().parent))
from jd_skill_extractor import (
    run_extraction,
    load_sponsors,
    to_slug,
    COGNITIVE_TIER,
)

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------
REPO_ROOT       = Path(__file__).resolve().parents[2]
CSV_PATH        = REPO_ROOT / "data" / "80-days-to-stay" / "80-days-csv" / "mapped_student_employment_targets_v3.csv"
OUTPUT_DIR      = REPO_ROOT / "data" / "skill-demand"
OUTPUT_XLSX     = OUTPUT_DIR / "skill_gap_report.xlsx"
DEFAULT_TARGETS = REPO_ROOT / "my_targets.txt"

DEFAULT_MIN_APPROVALS = 50

TIER_LABEL = {
    "HIGH": "HIGH — automation-resistant",
    "MED":  "MED — partially augmented",
    "LOW":  "LOW — execution risk",
}

# ---------------------------------------------------------------------------
# Colors for Excel
# ---------------------------------------------------------------------------
COLOR_HEADER     = "1F3864"   # dark navy
COLOR_HIGH       = "C6EFCE"   # green
COLOR_MED        = "FFEB9C"   # yellow
COLOR_LOW        = "FFC7CE"   # red
COLOR_SUBHEADER  = "D9E1F2"   # light blue
COLOR_WHITE      = "FFFFFF"
FONT_WHITE       = Font(name="Arial", bold=True, color="FFFFFF", size=11)
FONT_BOLD        = Font(name="Arial", bold=True, size=10)
FONT_NORMAL      = Font(name="Arial", size=10)

# ---------------------------------------------------------------------------
# CSV lookup (H1B + funding data only — extraction is in jd-skill-extractor.py)
# ---------------------------------------------------------------------------

def load_csv() -> pd.DataFrame:
    df = pd.read_csv(CSV_PATH, low_memory=False)
    df["Total Approvals"] = pd.to_numeric(df["Total Approvals"], errors="coerce").fillna(0)
    df["Total Denials"]   = pd.to_numeric(df["Total Denials"], errors="coerce").fillna(0)
    df["Approval_Rate"]   = pd.to_numeric(df["Approval_Rate"], errors="coerce")
    return df


def lookup_company(name: str, df: pd.DataFrame) -> dict:
    """Fuzzy match company name in 80-days CSV."""
    name_upper = name.upper().strip()
    slug = to_slug(name)

    # exact match first
    match = df[df["company_name"].str.upper().str.strip() == name_upper]

    # fallback: slug match
    if match.empty:
        match = df[df["company_name"].apply(to_slug) == slug]

    # fallback: partial match
    if match.empty:
        match = df[df["company_name"].str.upper().str.contains(
            re.escape(name_upper[:6]), na=False
        )]

    if match.empty:
        return {"found_in_csv": False}

    row = match.iloc[0]
    approvals = int(row.get("Total Approvals", 0))
    denials   = int(row.get("Total Denials", 0))
    total     = approvals + denials
    rate      = f"{row['Approval_Rate']:.1f}%" if pd.notna(row.get("Approval_Rate")) else (
                f"{(approvals/total*100):.1f}%" if total > 0 else "N/A")

    return {
        "found_in_csv":         True,
        "csv_name":             str(row.get("company_name", "")),
        "h1b_approvals":        approvals,
        "h1b_denials":          denials,
        "approval_rate":        rate,
        "funding_stage":        str(row.get("latest_funding_stage", "N/A") or "N/A"),
        "funding_date":         str(row.get("latest_funding_date",  "N/A") or "N/A"),
        "median_salary":        str(row.get("median_salary_offered","N/A") or "N/A"),
        "top_titles_sponsored": str(row.get("top_job_titles_sponsored", "") or ""),
    }


# ---------------------------------------------------------------------------
# Main scan
# ---------------------------------------------------------------------------

def scan(companies: list[str], df: pd.DataFrame, dry_run: bool) -> dict:
    """
    Orchestrates the full scan.
    Calls jd-skill-extractor.py for live JD extraction (titles + descriptions),
    then enriches each company with H-1B and funding data from the 80-days CSV.
    """
    print(f"  Calling jd-skill-extractor for live JD extraction ...")
    extraction = run_extraction(companies, dry_run=dry_run, verbose=True)

    # Enrich with H-1B + funding data from CSV for the scorecard
    scorecard: list[dict] = []
    for entry in extraction["scan_log"]:
        company  = entry["company"]
        csv_data = lookup_company(company, df)
        scorecard.append({
            "Company":          company,
            "In 80-Days CSV":   "Yes" if csv_data.get("found_in_csv") else "No",
            "H1B Approvals":    csv_data.get("h1b_approvals", "N/A"),
            "H1B Denials":      csv_data.get("h1b_denials",   "N/A"),
            "Approval Rate":    csv_data.get("approval_rate", "N/A"),
            "Funding Stage":    csv_data.get("funding_stage", "N/A"),
            "Funding Date":     csv_data.get("funding_date",  "N/A"),
            "Median Salary":    csv_data.get("median_salary", "N/A"),
            "ATS Detected":     entry["ats"].capitalize() if entry["ats"] != "not_found" else "Not Found",
            "Live DE/DA Roles": entry["de_da_jobs"],
            "Skills Found":     ", ".join(entry["skills_found"]) or "None",
        })

    return {
        "skill_counts":    extraction["skill_counts"],
        "skill_companies": extraction["skill_companies"],
        "scorecard":       scorecard,
        "live_titles":     extraction["job_records"],
        "not_found":       extraction["not_found"],
    }


# ---------------------------------------------------------------------------
# Excel builder
# ---------------------------------------------------------------------------

def style_header_row(ws, row: int, num_cols: int, color: str = COLOR_HEADER) -> None:
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font      = FONT_WHITE
        cell.fill      = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def autofit(ws, min_w: int = 10, max_w: int = 45) -> None:
    for col_cells in ws.columns:
        length = max(len(str(c.value or "")) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = max(min_w, min(length + 2, max_w))


def write_skill_rankings(wb: Workbook, skill_counts: dict, skill_companies: dict) -> None:
    ws = wb.create_sheet("Skill Rankings")
    headers = ["Rank", "Skill", "Appearances in Live Titles", "Companies", "Cognitive Tier", "Recommendation"]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    sorted_skills = sorted(skill_counts.items(), key=lambda x: x[1], reverse=True)
    tier_color = {"HIGH": COLOR_HIGH, "MED": COLOR_MED, "LOW": COLOR_LOW}
    reco = {
        "HIGH": "Prioritize — high demand + automation-resistant",
        "MED":  "Strengthen — useful but partially augmented by AI",
        "LOW":  "Deprioritize — execution work, high automation risk",
    }

    for rank, (skill, count) in enumerate(sorted_skills, 1):
        tier  = COGNITIVE_TIER.get(skill, "MED")
        cos   = len(skill_companies.get(skill, []))
        row   = [rank, skill, count, cos, TIER_LABEL.get(tier, tier), reco.get(tier, "")]
        ws.append(row)
        fill = PatternFill("solid", fgColor=tier_color.get(tier, COLOR_WHITE))
        for col in range(1, len(row) + 1):
            cell      = ws.cell(row=rank + 1, column=col)
            cell.font = FONT_NORMAL
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"
    autofit(ws)


def write_scorecard(wb: Workbook, scorecard: list[dict]) -> None:
    ws = wb.create_sheet("Sponsor Scorecard")
    if not scorecard:
        ws.append(["No data"])
        return
    headers = list(scorecard[0].keys())
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    for row_data in scorecard:
        row = list(row_data.values())
        ws.append(row)
        r = ws.max_row
        # Color rows by live role count
        live = row_data.get("Live DE/DA Roles", 0)
        if isinstance(live, int) and live > 0:
            fill = PatternFill("solid", fgColor=COLOR_HIGH)
        elif row_data.get("ATS Detected") == "Not Found":
            fill = PatternFill("solid", fgColor=COLOR_LOW)
        else:
            fill = PatternFill("solid", fgColor=COLOR_WHITE)
        for col in range(1, len(row) + 1):
            cell      = ws.cell(row=r, column=col)
            cell.font = FONT_NORMAL
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    autofit(ws)


def write_live_titles(wb: Workbook, live_titles: list[dict]) -> None:
    ws = wb.create_sheet("Live Job Titles")
    if not live_titles:
        ws.append(["No live DE/DA titles found across scanned companies."])
        return
    headers = list(live_titles[0].keys())
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    tier_color = {"HIGH": COLOR_HIGH, "MED": COLOR_MED, "LOW": COLOR_LOW}
    for row_data in live_titles:
        row = list(row_data.values())
        ws.append(row)
        r    = ws.max_row
        tier = row_data.get("Cognitive Tier", "MED")
        fill = PatternFill("solid", fgColor=tier_color.get(tier, COLOR_WHITE))
        for col in range(1, len(row) + 1):
            cell      = ws.cell(row=r, column=col)
            cell.font = FONT_NORMAL
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    autofit(ws)


def write_not_found(wb: Workbook, not_found: list[dict]) -> None:
    ws = wb.create_sheet("Not Found")
    ws.append(["Companies where no Greenhouse or Lever board was detected"])
    ws["A1"].font = Font(name="Arial", bold=True, size=11)
    ws.append([])

    if not not_found:
        ws.append(["All companies had a detectable ATS board."])
        return

    headers = list(not_found[0].keys())
    ws.append(headers)
    style_header_row(ws, 3, len(headers))

    for row_data in not_found:
        row = list(row_data.values())
        ws.append(row)
        r = ws.max_row
        for col in range(1, len(row) + 1):
            cell      = ws.cell(row=r, column=col)
            cell.font = FONT_NORMAL
            cell.fill = PatternFill("solid", fgColor=COLOR_LOW)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    autofit(ws)


def build_excel(results: dict, run_ts: str, dry_run: bool) -> None:
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    write_skill_rankings(wb, results["skill_counts"], results["skill_companies"])
    write_scorecard(wb, results["scorecard"])
    write_live_titles(wb, results["live_titles"])
    write_not_found(wb, results["not_found"])

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_XLSX)
    print(f"\nExcel report saved to: {OUTPUT_XLSX}")


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(
        description="Scan live DE/DA skill demand at verified H-1B sponsor companies."
    )
    parser.add_argument("--targets",      default=str(DEFAULT_TARGETS),
                        help="Path to company names file (one per line).")
    parser.add_argument("--dry-run",      action="store_true",
                        help="Skip API calls — show what would be scanned.")
    parser.add_argument("--all-sponsors", action="store_true",
                        help="Ignore targets file — scan all 38 verified DE/DA sponsors.")
    parser.add_argument("--min-approvals", type=int, default=DEFAULT_MIN_APPROVALS,
                        help="Minimum H-1B approvals when using --all-sponsors.")
    args = parser.parse_args()

    run_ts = datetime.now(timezone.utc).isoformat()
    print(f"\ncase-de-da-live-skill-gap | {run_ts}")
    print("=" * 60)

    # Load 80-days CSV
    print(f"\n[1/3] Loading 80-days CSV ...")
    if not CSV_PATH.exists():
        print(f"ERROR: CSV not found at {CSV_PATH}")
        sys.exit(1)
    df = load_csv()
    print(f"      Loaded {len(df):,} companies")

    # Resolve company list
    if args.all_sponsors:
        de_da_mask = (
            (df["Total Approvals"] >= args.min_approvals) &
            (df["top_job_titles_sponsored"].fillna("").str.contains(
                r"data\s+engineer|data\s+analyst|analytics\s+engineer",
                case=False, regex=True, na=False
            ))
        )
        companies = df[de_da_mask]["company_name"].tolist()
        print(f"      --all-sponsors: {len(companies)} verified DE/DA sponsors")
    else:
        targets_path = Path(args.targets)
        if not targets_path.exists():
            print(f"\nERROR: Targets file not found: {targets_path}")
            print("Create my_targets.txt in the repo root with one company name per line.")
            print("\nExample:")
            print("  Databricks")
            print("  CVS Health")
            print("  Experian")
            sys.exit(1)
        companies = [l.strip() for l in targets_path.read_text().splitlines() if l.strip()]
        print(f"      Loaded {len(companies)} companies from {targets_path.name}")

    if not companies:
        print("ERROR: No companies to scan.")
        sys.exit(1)

    # Scan
    print(f"\n[2/3] Scanning {'(DRY RUN — no API calls)' if args.dry_run else '(LIVE)'}...")
    results = scan(companies, df, dry_run=args.dry_run)

    # Build Excel
    print(f"\n[3/3] Building Excel report ...")
    build_excel(results, run_ts, dry_run=args.dry_run)

    # Terminal summary
    if results["skill_counts"] and not args.dry_run:
        print("\n--- Top 10 Skills (preview) ---")
        for i, (skill, count) in enumerate(
            sorted(results["skill_counts"].items(), key=lambda x: x[1], reverse=True)[:10], 1
        ):
            tier  = COGNITIVE_TIER.get(skill, "?")
            cos_n = len(results["skill_companies"].get(skill, []))
            print(f"  {i:2}. {skill:<12} {count:3} appearances  {cos_n:2} companies  [{tier}]")
    elif args.dry_run:
        print(f"\nDRY RUN complete — {len(companies)} companies would be scanned")
        print(f"Output would be written to: {OUTPUT_XLSX}")

    print("\nDone.")


if __name__ == "__main__":
    main()